"""
A안 — 스킬 미사용. python-pptx 기본 기능만으로 선행연구 조사 발표 자료를 만든다.
실행: python3 make_ppt.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

OUT = Path(__file__).parent / "선행연구조사_게임닉네임.pptx"

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x3B, 0x73)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Cm(33.87)   # 16:9
prs.slide_height = Cm(19.05)

BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def set_run(run, size, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color if color else RGBColor(0, 0, 0)


def title_slide(title, subtitle, author):
    slide = add_slide()
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY

    tf = textbox(slide, 3, 6.5, 28, 4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), 40, True, WHITE)
    p.runs[0].text = title

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(18)
    r2 = p2.add_run()
    r2.text = subtitle
    set_run(r2, 18, False, RGBColor(0xCC, 0xD4, 0xE8))

    tf3 = textbox(slide, 3, 15, 28, 2)
    p3 = tf3.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = author
    set_run(r3, 14, False, RGBColor(0xCC, 0xD4, 0xE8))


def content_slide(title, bullets, size=18):
    """bullets: (텍스트, 들여쓰기 단계) 목록. size 는 1단 글자 크기."""
    slide = add_slide()

    tf = textbox(slide, 2, 1.4, 30, 2)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, 28, True, NAVY)

    line = slide.shapes.add_shape(1, Cm(2), Cm(3.5), Cm(29.9), Cm(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = NAVY
    line.line.fill.background()

    tf2 = textbox(slide, 2.4, 4.4, 29, 13)
    first = True
    for text, level in bullets:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(10)
        r = p.add_run()
        if level == 0:
            r.text = "• " + text
            set_run(r, size, False)
        else:
            r.text = "- " + text
            set_run(r, size - 3, False, GRAY)


def table_slide(title, headers, rows):
    slide = add_slide()

    tf = textbox(slide, 2, 1.4, 30, 2)
    r = tf.paragraphs[0].add_run()
    r.text = title
    set_run(r, 28, True, NAVY)

    line = slide.shapes.add_shape(1, Cm(2), Cm(3.5), Cm(29.9), Cm(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = NAVY
    line.line.fill.background()

    n_rows, n_cols = len(rows) + 1, len(headers)
    shape = slide.shapes.add_table(n_rows, n_cols, Cm(2), Cm(4.4), Cm(29.9), Cm(2))
    table = shape.table

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        para = cell.text_frame.paragraphs[0]
        set_run(para.runs[0], 13, True, WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
            para = cell.text_frame.paragraphs[0]
            set_run(para.runs[0], 11, False)


# ----------------------------------------------------------------------
# 1. 표지
title_slide(
    "게임 닉네임의 형태·의미 분석",
    "선행연구 조사",
    "AI기반콘텐츠창업워크숍2 · 2026.10.14",
)

# 2. 연구 배경과 목적
content_slide(
    "1. 연구 배경과 목적",
    [
        ("게임 닉네임은 사용자가 스스로 만든 언어 자료다", 0),
        ("실명과 달리 제약이 적어 조어·변형·차용이 자유롭게 나타난다", 1),
        ("따라서 형태론·의미론 연구의 자료로서 가치가 있다", 1),
        ("기존 연구는 언어별·소규모 수작업 분석에 머물러 있다", 0),
        ("수천 건 규모의 분류를 사람이 직접 수행해 왔다", 1),
        ("언어가 바뀌면 분석 틀을 다시 만들어야 한다", 1),
        ("본 연구의 목표", 0),
        ("선행연구의 분류 체계를 LLM 기반 시스템으로 구현한다", 1),
        ("다국어로 확장해 언어 간 비교를 가능하게 한다", 1),
        ("최종적으로 작명 동기와 행위의 관계를 본다", 1),
    ],
)

# 3. 선행연구 지도
content_slide(
    "2. 선행연구 지도 — 두 갈래",
    [
        ("갈래 A — Semantic network 방법론 (2편)", 0),
        ("키워드 동시출현으로 의미 구조를 드러내는 분석 절차", 1),
        ("'무엇으로 분석할 것인가'에 해당한다", 1),
        ("갈래 B — 닉네임의 형태·통사·의미 분석 (5편)", 0),
        ("사용자명을 언어학적으로 분류한 실증 연구들", 1),
        ("'무엇을 분석할 것인가'에 해당한다", 1),
        ("두 갈래는 아직 만나지 않았다", 0),
        ("닉네임 연구는 수작업 분류에, 네트워크 방법론은 일반 코퍼스에 머물러 있다", 1),
        ("이 지점이 본 연구의 출발점이다", 1),
    ],
)

# 4. 갈래 A
content_slide(
    "3. 갈래 A — Semantic network 방법론",
    [
        ("Corpus Linguistics, Network Analysis and Co-occurrence Matrices", 0),
        ("International Journal of English Studies, Univ. of Murcia", 1),
        ("빈도 → 키워드 → 연어 분석 → 동시출현 행렬 → semantic network 의 절차", 1),
        ("코퍼스에서 네트워크까지 가는 경로를 단계로 정리했다", 1),
        ("Omodei, Cointet & Poibeau (2014), LREC'14, pp. 2972-2978", 0),
        ("언어 패턴과 통계를 혼합한 키워드 추출", 1),
        ("통계만 쓰는 방식보다 의미 있는 키워드를 더 잘 잡는다", 1),
        ("추출한 키워드의 동시출현으로 네트워크를 구축한다", 1),
        ("본 연구와의 접점", 0),
        ("닉네임 구성 요소를 키워드로 보면 같은 절차를 적용할 수 있다", 1),
    ],
)

# 5. 갈래 B
content_slide(
    "4. 갈래 B — 닉네임의 형태·의미 분석",
    [
        ("GD Usernames and Euphemism: A Morphosemantic Analysis", 0),
        ("게임 사이트 사용자명의 단어 형성, 철자 변형, 차용, 의미 혁신", 1),
        ("Szymanski (2013)", 0),
        ("닉네임 7,456개의 형태론 분석 및 문자 선택 동기", 1),
        ("Dubiec-Stach (2021)", 0),
        ("독일어·폴란드어 닉네임의 의미·형태·통사·문자 양식 분류", 1),
        ("두 언어를 나란히 놓은 비교 연구", 1),
        ("Kaziaba (2022)", 0),
        ("5,000개 사례의 다층위 분석, 온라인 게임 포함", 1),
        ("Fomin & Chobotar (2022)", 0),
        ("게임 내 닉네임의 작명 동기와 생성 기제", 1),
    ],
)

# 6. 비교표
table_slide(
    "5. 선행연구 비교",
    ["연구", "자료 규모", "언어", "분석 층위", "방법"],
    [
        ["Murcia (연도 미상)", "일반 코퍼스", "영어", "키워드·연어", "동시출현 행렬"],
        ["Omodei et al. (2014)", "학술 코퍼스", "영어", "키워드", "패턴+통계 혼합"],
        ["GD Usernames (연도 미상)", "미상", "영어", "형태·의미", "수작업 분류"],
        ["Szymanski (2013)", "7,456", "폴란드어", "형태", "수작업 분류"],
        ["Dubiec-Stach (2021)", "미상", "독일어·폴란드어", "의미·형태·통사·문자", "수작업 분류"],
        ["Kaziaba (2022)", "5,000", "러시아어", "다층위", "수작업 분류"],
        ["Fomin & Chobotar (2022)", "미상", "우크라이나어", "동기·생성 기제", "수작업 분류"],
    ],
)

# 7. 연구 공백
content_slide(
    "6. 연구 공백",
    [
        ("공백 1 — 분석 규모의 한계", 0),
        ("최대 7,456건. 모두 사람이 직접 분류했다", 1),
        ("게임 플랫폼 한 곳의 하루치 신규 계정에도 못 미치는 규모다", 1),
        ("공백 2 — 언어 간 비교의 단절", 0),
        ("폴란드어·독일어·러시아어·우크라이나어가 각각 따로 연구되었다", 1),
        ("분류 체계가 연구마다 달라 결과를 포개 볼 수 없다", 1),
        ("한국어·일본어 등 비유럽 언어는 거의 다뤄지지 않았다", 1),
        ("공백 3 — 동기와 행위의 분리", 0),
        ("Fomin & Chobotar가 작명 동기를 다루었으나 설문·면접에 의존한다", 1),
        ("닉네임과 실제 플레이 행위를 연결한 연구는 찾지 못했다", 1),
    ],
)

# 8. 제안 방향
content_slide(
    "7. 제안 연구 방향",
    [
        ("1단계 — 분류 체계의 코드화", 0),
        ("Dubiec-Stach의 네 층위(의미·형태·통사·문자)를 LLM 프롬프트로 구현한다", 1),
        ("선행연구가 수작업으로 한 분류를 자동화하는 것이 1차 목표다", 1),
        ("2단계 — 다국어 확장", 0),
        ("같은 분류 체계를 언어별로 적용해 비교 가능한 결과를 만든다", 1),
        ("한국어를 포함해 공백 2를 직접 겨냥한다", 1),
        ("3단계 — semantic network 결합", 0),
        ("닉네임 구성 요소를 키워드로 보고 동시출현 네트워크를 그린다", 1),
        ("갈래 A의 방법을 갈래 B의 자료에 적용하는 지점이다", 1),
        ("4단계 — 동기와 행위의 연결", 0),
        ("분류 결과를 플레이 데이터와 대조한다", 1),
    ],
)

# 9. 참고문헌
content_slide(
    "8. 참고문헌",
    [
        ("Corpus Linguistics, Network Analysis and Co-occurrence Matrices. "
         "International Journal of English Studies, Univ. of Murcia. [서지정보 보완 필요]", 0),
        ("Omodei, E., Cointet, J.-P., & Poibeau, T. (2014). LREC'14, pp. 2972-2978.", 0),
        ("GD Usernames and Euphemism: A Morphosemantic Analysis. [서지정보 보완 필요]", 0),
        ("Szymanski (2013). 닉네임 7,456개의 형태론 분석. [서지정보 보완 필요]", 0),
        ("Dubiec-Stach (2021). 독일어·폴란드어 닉네임 분류. [서지정보 보완 필요]", 0),
        ("Kaziaba (2022). 닉네임 5,000개의 다층위 분석. [서지정보 보완 필요]", 0),
        ("Fomin & Chobotar (2022). 게임 내 닉네임의 작명 동기. [서지정보 보완 필요]", 0),
        ("※ 저자 전체 이름·학술지명·권호·DOI 를 확인하지 못했다. "
         "검색 도구 없이 주어진 목록만 사용한 조건의 한계다.", 0),
    ],
    size=14,
)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"저장: {OUT}  (슬라이드 {len(prs.slides.__iter__.__self__._sldIdLst)}장)")
